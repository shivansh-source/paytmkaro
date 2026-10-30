"""The round-2 sample documents (bureau, Udyam, financials, sanction letter, settlement),
PDF table reconstruction, and the privacy rules that must hold on them."""

import re
from pathlib import Path

import pytest
from PIL import Image

from app.documents.fakes import FakeOcr
from app.documents.ingest import load_document
from app.documents.models import LoadStatus, OcrLine, SourceFormat, TextSource
from app.documents.privacy import verhoeff_valid

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "documents"

NEW_SAMPLES = [
    "bureau_report.pdf",
    "bureau_report_commercial.pdf",
    "bureau_report_scanned.jpg",
    "udyam_certificate.pdf",
    "udyam_certificate.jpg",
    "financial_statements.pdf",
    "financial_statements.xlsx",
    "sanction_letter.pdf",
    "sanction_letter.docx",
    "settlement_report.csv",
    "settlement_report.xlsx",
]


def load(name: str, **kwargs):
    return load_document((FIXTURES / name).read_bytes(), name, **kwargs)


def money(text: str) -> float:
    return float(re.sub(r"[^\d.]", "", text.replace("Rs.", "")))


@pytest.mark.parametrize("name", NEW_SAMPLES)
def test_new_samples_load_and_contain_no_aadhaar(name):
    ocr = FakeOcr([OcrLine(text="SAMPLE", confidence=0.9, bbox=(0, 0, 100, 20))])
    doc = load(name, ocr=ocr)
    assert doc.status is LoadStatus.OK, doc.message
    assert doc.pages
    assert "XXXX XXXX" not in doc.model_dump_json()  # nothing was masked: nothing to mask


def test_scanned_samples_are_images_without_native_text():
    for name in ("bureau_report_scanned.jpg", "udyam_certificate.jpg"):
        doc = load(name)
        assert doc.pages[0].text_source is TextSource.NONE and doc.pages[0].image_png


# ---------- bureau ----------


def test_consumer_bureau_report_keeps_other_lenders_data_for_the_extractor_to_drop():
    doc = load("bureau_report.pdf")
    page = doc.pages[0]
    assert "DEMO SCORE: 705" in page.text
    assert "9876543210" in page.text and "12, Market Road" in page.text
    accounts, enquiries = page.tables
    assert accounts[0][:3] == ["Lender", "Type", "Account No"]
    assert len(accounts) == 4 and accounts[2][2] == "LN/2023/778812" and accounts[2][6] == "30"
    assert accounts[3][2] == "4387000012345678"  # 16-digit card number is not an Aadhaar number
    assert len(enquiries) == 4 and enquiries[1][1] == "Demo Bank A"


def test_commercial_bureau_report_has_cmr_rank():
    page = load("bureau_report_commercial.pdf").pages[0]
    assert "CMR RANK: 4" in page.text
    assert page.tables[0][1][:2] == ["Demo Bank Z", "Cash Credit"]


# ---------- Udyam ----------


def test_udyam_certificate_text():
    text = load("udyam_certificate.pdf").pages[0].text
    for expected in ("UDYAM-MH-00-1234567", "Type of Enterprise: Small", "46900", "Date of Udyam Registration: 10/05/2022"):
        assert expected in text


# ---------- financial statements ----------


def test_financial_pdf_pages_and_tables_balance():
    doc = load("financial_statements.pdf")
    balance_sheet, profit_loss = doc.pages
    assert "(Rs. in Lakhs)" in balance_sheet.text and "UDIN: 25123456ABCDEF1234" in balance_sheet.text
    assert "Statement of Profit and Loss" in profit_loss.text
    rows = {r[0]: r[1:] for r in balance_sheet.tables[0]}
    assert rows["Particulars"] == ["31-Mar-2025", "31-Mar-2024"]
    assert rows["Total equity and liabilities"] == rows["Total assets"] == ["276.65", "240.00"]
    pl = {r[0]: r[1:] for r in profit_loss.tables[0]}
    assert pl["Profit after tax"] == ["27.00", "18.75"]
    assert money(pl["Total income"][0]) - money(pl["Total expenses"][0]) == money(pl["Profit before tax"][0])


def test_financial_xlsx_is_in_rupees():
    doc = load("financial_statements.xlsx")
    assert [p.index for p in doc.pages] == [0, 1]
    rows = {r[0]: r[1:] for r in doc.pages[0].tables[0] if len(r) == 3}
    assert rows["Total assets"] == ["27665000", "24000000"]
    assert "(Amount in Rs.)" in doc.pages[0].text and "Provisional" in doc.pages[0].text
    pl = {r[0]: r[1:] for r in doc.pages[1].tables[0] if len(r) == 3}
    assert pl["Profit after tax"] == ["2700000", "1875000"]


# ---------- sanction letter ----------


def sanction_values(table) -> dict[str, str]:
    return {row[0]: row[1] for row in table if len(row) == 2}


@pytest.mark.parametrize("name", ["sanction_letter.pdf", "sanction_letter.docx"])
def test_sanction_letter_numbers_are_mathematically_consistent(name):
    page = load(name).pages[0]
    values = sanction_values(page.tables[-1])
    principal, months = money(values["Sanctioned loan amount"]), int(values["Number of instalments"])
    rate = float(re.search(r"([\d.]+)%", values["Rate of interest"]).group(1)) / 1200
    emi = principal * rate * (1 + rate) ** months / ((1 + rate) ** months - 1)
    assert money(values["EMI"]) == pytest.approx(emi, abs=0.01)
    assert money(values["Total interest charge"]) == pytest.approx(money(values["EMI"]) * months - principal, abs=0.01)
    fees = money(values["Processing fee"].split("(")[0]) + money(values["GST on processing fee"].split("(")[0])
    assert money(values["Net amount disbursed"]) == pytest.approx(principal - fees, abs=0.01)
    assert money(values["Total amount to be paid"]) == pytest.approx(money(values["EMI"]) * months + fees, abs=0.01)
    assert "Cooling-off period" in values and "KEY FACT STATEMENT" in page.text


def test_sanction_apr_is_the_irr_of_the_borrowers_cash_flows():
    values = sanction_values(load("sanction_letter.pdf").pages[0].tables[-1])
    emi, net, months = money(values["EMI"]), money(values["Net amount disbursed"]), int(values["Number of instalments"])
    monthly = float(values["Annual Percentage Rate (APR)"].rstrip("%")) / 1200
    assert sum(emi / (1 + monthly) ** k for k in range(1, months + 1)) == pytest.approx(net, rel=1e-3)


# ---------- settlement report ----------


@pytest.mark.parametrize("name", ["settlement_report.csv", "settlement_report.xlsx"])
def test_settlement_report_rows_and_unmasked_references(name):
    doc = load(name)
    rows = doc.pages[0].tables[0]
    header, body = rows[0], rows[1:]
    assert header[0] == "transaction_id" and "payer_upi_id" in header and "payer_phone" in header
    assert len(body) == 60
    ids = [r[0] for r in body]
    # Premise: every id has the shape of an Aadhaar number (12 digits, first digit 2-9, valid check digit)...
    assert all(re.fullmatch(r"[2-9]\d{11}", i) and verhoeff_valid(i) for i in ids)
    # ...and still none is masked, because nothing on the page mentions Aadhaar.
    assert not any("XXXX" in cell for r in body for cell in r)
    assert all(i in doc.pages[0].text for i in ids)
    gross, fee, net = (header.index(c) for c in ("gross_amount", "fee", "net_amount"))
    assert all(money(r[gross]) - money(r[fee]) == pytest.approx(money(r[net]), abs=0.01) for r in body)
    assert body[0][header.index("date")] >= "2025-04-01" and body[-1][header.index("date")] <= "2025-06-30"


def test_settlement_ids_ARE_masked_when_the_sheet_mentions_aadhaar():
    # The context rule from the other side: same shape of id, but an Aadhaar keyword on the page.
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["customer Aadhaar", "reference"])
    sheet.append([int(load("settlement_report.csv").pages[0].tables[0][1][0]), "x"])
    out = io.BytesIO()
    workbook.save(out)
    doc = load_document(out.getvalue(), "mixed.xlsx")
    assert doc.pages[0].tables[0][1][0].startswith("XXXX XXXX ")


# ---------- PDF table reconstruction ----------


def test_bank_statement_pdf_table_keeps_empty_cells_in_their_columns():
    [table] = load("bank_statement.pdf").pages[0].tables
    assert table[0] == ["Date", "Narration", "Chq/Ref No", "Withdrawal", "Deposit", "Balance"]
    assert len(table) == 8 and all(len(row) == 6 for row in table)
    assert table[1] == ["02-04-2025", "UPI/412345678901/Payment from Sharma Stores", "REF1000", "", "25000.00", "175000.00"]
    assert table[2][3:] == ["18000.00", "", "157000.00"]


def test_pdf_prose_does_not_become_a_table():
    assert load("gst_reg06.pdf").pages[0].tables == []


def test_scanned_pages_get_no_tables():
    doc = load("gst_reg06_scanned.pdf", ocr=FakeOcr([OcrLine(text="x y", confidence=0.9, bbox=(0, 0, 9, 9))]))
    assert doc.pages[0].tables == []


# ---------- privacy rules still hold after the privacy.py change ----------

A, B, C = "2345", "6789", "0124"


def test_year_followed_by_newline_and_eight_digits_is_not_an_aadhaar_number():
    text = "Valid till 2025\n6789 0123"
    doc = load_document(text.encode(), "note.txt")
    assert doc.pages[0].text == text


def test_year_on_one_ocr_row_and_eight_digits_on_the_next_is_not_masked():
    lines = [
        OcrLine(text="Valid till 2025", confidence=0.9, bbox=(0, 0, 300, 40)),
        OcrLine(text="6789 0123", confidence=0.9, bbox=(0, 100, 220, 140)),
    ]
    import io

    image = io.BytesIO()
    Image.new("RGB", (400, 200), "white").save(image, format="PNG")
    page = load_document(image.getvalue(), "a.png", ocr=FakeOcr(lines)).pages[0]
    assert page.text == "Valid till 2025\n6789 0123"


@pytest.mark.parametrize(
    "number",
    [f"{A} {B} {C}", f"{A}  {B}  {C}", f"{A}\t{B}\t{C}", f"{A} {B} {C}", f"{A} - {B} - {C}"],
)
def test_odd_separators_are_masked_in_text_and_cells(number):
    doc = load_document(f"Aadhaar: {number}".encode(), "n.txt")
    assert doc.pages[0].text == f"Aadhaar: XXXX XXXX {C}"
    csv = load_document(f"name,number\nR. Kumar,{number}\nx,y\n".encode(), "n.csv")
    assert csv.pages[0].tables[0][1] == ["R. Kumar", f"XXXX XXXX {C}"]
