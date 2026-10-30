"""DOCX, XLSX, XLS, CSV and TXT ingest."""

import datetime as dt
import io
from pathlib import Path

import docx
import openpyxl
import pytest
from PIL import Image

from app.documents.fakes import FakeOcr
from app.documents.ingest import load_document
from app.documents.ingest.sheets import format_cell
from app.documents.models import LoadStatus, OcrLine, SourceFormat, TextSource
from app.documents.privacy import verhoeff_valid

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "documents"
AADHAAR = next(f"23456789012{d}" for d in "0123456789" if verhoeff_valid(f"23456789012{d}"))
OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def save(document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ---------- DOCX ----------


def test_docx_body_tables_and_embedded_photo_ocr():
    ocr = FakeOcr([OcrLine(text="INCOME TAX DEPARTMENT ABCDE1234F", confidence=0.9, bbox=(0, 0, 300, 30))])
    doc = load_document((FIXTURES / "kyc_bundle.docx").read_bytes(), "kyc.pdf", ocr=ocr)
    assert doc.status is LoadStatus.OK and doc.source_format is SourceFormat.DOCX
    body, photo = doc.pages
    assert body.text_source is TextSource.NATIVE
    assert "KYC documents" in body.text and "Proprietorship" in body.text
    assert ["PAN", "ABCDE1234F"] in body.tables[0]
    assert body.image_png is None
    assert photo.index == 1 and photo.text_source is TextSource.OCR and photo.image_png
    assert ocr.calls == 1  # python-docx's own docProps thumbnail is not a page


def test_docx_masks_aadhaar_in_paragraphs_and_tables():
    document = docx.Document()
    document.add_paragraph(f"Aadhaar of proprietor: {AADHAAR[:4]} {AADHAAR[4:8]} {AADHAAR[8:]}")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Aadhaar No", AADHAAR
    doc = load_document(save(document), "kyc.docx")
    page = doc.pages[0]
    assert AADHAAR[:8] not in page.text
    assert page.tables[0][0] == ["Aadhaar No", f"XXXX XXXX {AADHAAR[-4:]}"]
    assert AADHAAR[:8] not in doc.model_dump_json()


def test_docx_tiny_images_are_skipped():
    document = docx.Document()
    document.add_paragraph("Letterhead")
    logo = io.BytesIO()
    Image.new("RGB", (32, 32), "red").save(logo, format="PNG")
    document.add_picture(logo)
    doc = load_document(save(document), "letter.docx", ocr=FakeOcr())
    assert len(doc.pages) == 1


# ---------- XLSX / XLS ----------


def test_xlsx_statement_one_page_per_sheet_with_clean_values():
    doc = load_document((FIXTURES / "bank_statement.xlsx").read_bytes(), "statement.csv")
    assert doc.status is LoadStatus.OK and doc.source_format is SourceFormat.XLSX
    statement, summary = doc.pages
    rows = statement.tables[0]
    header = rows.index(["Date", "Narration", "Chq/Ref No", "Withdrawal", "Deposit", "Balance"])
    first = rows[header + 1]
    assert first[0] == "2025-04-02"
    assert first[4] == "25000" and first[5] == "175000"
    assert any(row[4] == "1500000" for row in rows[header + 1 :])
    assert any(row[3] == "31250.25" for row in rows[header + 1 :])
    assert "UPI/412345678901/Payment from Sharma Stores" in statement.text  # UPI refs survive
    assert "\t" in statement.text
    assert summary.tables == [[["Total credits", "1567500.5"]]]
    assert all(p.image_png is None for p in doc.pages)


def test_xlsx_masks_numeric_aadhaar_cell_when_sheet_mentions_aadhaar():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Aadhaar"])
    sheet.append(["R. Kumar", int(AADHAAR)])
    sheet.append(["UPI ref", 412345678901])
    doc = load_document(save(workbook), "kyc.xlsx")
    rows = doc.pages[0].tables[0]
    assert rows[1] == ["R. Kumar", f"XXXX XXXX {AADHAAR[-4:]}"]
    assert rows[2] == ["UPI ref", "412345678901"]  # fails Verhoeff: not an Aadhaar number
    assert AADHAAR not in doc.model_dump_json()


def test_xlsx_sheet_limit(monkeypatch):
    from app import config

    monkeypatch.setattr(config, "MAX_PAGES", 1)
    doc = load_document((FIXTURES / "bank_statement.xlsx").read_bytes(), "s.xlsx")
    assert doc.status is LoadStatus.TOO_LARGE


def test_empty_workbook_still_gives_a_page():
    doc = load_document(save(openpyxl.Workbook()), "empty.xlsx")
    assert doc.status is LoadStatus.OK
    assert len(doc.pages) == 1 and doc.pages[0].text == ""
    assert any("No text" in w for w in doc.warnings)


@pytest.mark.parametrize("name", ["statement.xlsx", "kyc.docx", "file.bin"])
def test_encrypted_office_file_needs_password(name):
    data = OLE2 + b"\x00" * 504 + "EncryptionInfo".encode("utf-16-le") + "EncryptedPackage".encode("utf-16-le")
    doc = load_document(data, name)
    assert doc.status is LoadStatus.PASSWORD_REQUIRED
    assert "password" in doc.message.lower()


def test_xls_is_read_through_xlrd(monkeypatch):
    import xlrd

    class Cell:
        def __init__(self, ctype, value):
            self.ctype, self.value = ctype, value

    class Sheet:
        nrows = 2

        def row(self, r):
            return [
                [Cell(xlrd.XL_CELL_TEXT, "Date"), Cell(xlrd.XL_CELL_TEXT, "Credit")],
                [Cell(xlrd.XL_CELL_DATE, 45748.0), Cell(xlrd.XL_CELL_NUMBER, 1500000.0)],
            ][r]

    class Book:
        nsheets, datemode = 1, 0

        def sheet_by_index(self, i):
            return Sheet()

        def release_resources(self):
            pass

    monkeypatch.setattr(xlrd, "open_workbook", lambda **kwargs: Book())
    data = OLE2 + b"\x00" * 504 + "Workbook".encode("utf-16-le")
    doc = load_document(data, "statement.xls")
    assert doc.source_format is SourceFormat.XLS and doc.status is LoadStatus.OK
    assert doc.pages[0].tables == [[["Date", "Credit"], ["2025-04-01", "1500000"]]]


def test_encrypted_xls(monkeypatch):
    import xlrd

    def encrypted(**kwargs):
        raise xlrd.XLRDError("Workbook is encrypted")

    monkeypatch.setattr(xlrd, "open_workbook", encrypted)
    doc = load_document(OLE2 + b"\x00" * 504 + "Workbook".encode("utf-16-le"), "s.xls")
    assert doc.status is LoadStatus.PASSWORD_REQUIRED


@pytest.mark.parametrize(
    "value,expected",
    [
        (1500000.0, "1500000"),
        (0.1 + 0.2, "0.3"),
        (1234.5, "1234.5"),
        (12345678901.25, "12345678901.25"),
        (dt.datetime(2025, 4, 1), "2025-04-01"),
        (dt.datetime(2025, 4, 1, 10, 30), "2025-04-01 10:30:00"),
        (dt.date(2025, 4, 1), "2025-04-01"),
        (True, "TRUE"),
        (None, ""),
        (42, "42"),
        ("  text ", "text"),
    ],
)
def test_cell_formatting(value, expected):
    assert format_cell(value) == expected


# ---------- CSV / TXT ----------


def test_csv_statement():
    doc = load_document((FIXTURES / "bank_statement.csv").read_bytes(), "statement.xlsx")
    assert doc.source_format is SourceFormat.CSV
    [page] = doc.pages
    assert page.tables[0][0] == ["Date", "Narration", "Chq/Ref No", "Withdrawal", "Deposit", "Balance"]
    assert page.tables[0][1][1] == "UPI/412345678901/Payment from Sharma Stores"
    assert page.text_source is TextSource.NATIVE


def test_csv_semicolon_and_cp1252():
    data = "Date;Narration;Amount\n01-04-2025;Café rent;1500\n02-04-2025;UPI;200\n".encode("cp1252")
    doc = load_document(data, "s.csv")
    assert doc.pages[0].tables[0][1] == ["01-04-2025", "Café rent", "1500"]


def test_utf8_bom_csv():
    data = "﻿Date,Amount\n01-04-2025,100\n02-04-2025,200\n".encode("utf-8")
    doc = load_document(data, "s.csv")
    assert doc.pages[0].tables[0][0] == ["Date", "Amount"]


def test_txt_is_masked():
    data = f"Aadhaar: {AADHAAR[:4]}-{AADHAAR[4:8]}-{AADHAAR[8:]}\nName: R. Kumar".encode()
    doc = load_document(data, "note.txt")
    assert doc.source_format is SourceFormat.TXT
    assert doc.pages[0].text == f"Aadhaar: XXXX XXXX {AADHAAR[-4:]}\nName: R. Kumar"


def test_txt_with_wide_or_non_breaking_spaces_is_masked():
    a, b, c = AADHAAR[:4], AADHAAR[4:8], AADHAAR[8:]
    doc = load_document(f"No. {a}   {b}   {c}\nNo. {a} {b} {c}".encode(), "note.txt")
    assert doc.pages[0].text == f"No. XXXX XXXX {c}\nNo. XXXX XXXX {c}"
